"""Integration tests for Excel workbook generation, sheet count, formatting, and data integrity."""

import pytest
from pathlib import Path
import openpyxl

from src.config import EXCEL_OUTPUT_FILE

EXPECTED_SHEETS = [
    "Executive Dashboard",
    "Monthly Trends",
    "Category Analysis",
    "Product Analysis",
    "Regional Analysis",
    "Discount & Profitability",
    "Customer Analysis",
    "Cleaned Data",
    "Data Dictionary",
    "Methodology & KPI Definitions",
]


def test_excel_file_exists():
    """Verify that the generated Excel workbook exists."""
    assert EXCEL_OUTPUT_FILE.exists(), f"Excel workbook missing at: {EXCEL_OUTPUT_FILE}"
    assert EXCEL_OUTPUT_FILE.stat().st_size > 50000, "Excel file is unusually small"


def test_excel_ten_sheets_exist():
    """Verify all 10 expected business sheets are present in the workbook."""
    wb = openpyxl.load_workbook(EXCEL_OUTPUT_FILE, data_only=False)
    sheet_names = wb.sheetnames
    assert len(sheet_names) == 10, f"Expected 10 sheets, found {len(sheet_names)}: {sheet_names}"
    for expected in EXPECTED_SHEETS:
        assert expected in sheet_names, f"Missing required sheet: {expected}"


def test_executive_dashboard_kpis_populated():
    """Verify Executive Dashboard has KPI cards and non-empty values."""
    wb = openpyxl.load_workbook(EXCEL_OUTPUT_FILE, data_only=True)
    ws = wb["Executive Dashboard"]

    # Title Banner Check
    title_val = ws.cell(row=1, column=1).value
    assert "RETAIL SALES" in str(title_val).upper()

    # KPI labels in Row 4, KPI values in Row 5
    rev_label = ws.cell(row=4, column=1).value
    assert "REVENUE" in str(rev_label).upper()
    rev_val = ws.cell(row=5, column=1).value
    assert rev_val is not None and float(rev_val) > 1000000.0  # Total revenue is ~$2.29M

    # Check Net Profit card
    profit_label = ws.cell(row=4, column=2).value
    assert "PROFIT" in str(profit_label).upper()
    profit_val = ws.cell(row=5, column=2).value
    assert profit_val is not None and float(profit_val) > 200000.0  # Net profit is ~$286K


def test_cleaned_data_sheet_structure():
    """Verify Cleaned Data sheet contains headers and rows with auto-filter enabled."""
    wb = openpyxl.load_workbook(EXCEL_OUTPUT_FILE, data_only=True)
    ws = wb["Cleaned Data"]

    # Header in row 4
    header_order_id = ws.cell(row=4, column=1).value
    assert "Order Id" in str(header_order_id) or "ORDER ID" in str(header_order_id).upper()

    # Data row 5
    first_order_val = ws.cell(row=5, column=1).value
    assert first_order_val is not None
    assert str(first_order_val).startswith("CA-") or str(first_order_val).startswith("US-")

    # Check auto-filter
    assert ws.auto_filter.ref is not None
